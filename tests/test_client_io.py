import io
from datetime import timedelta

import asyncpg
import pytest
from openpyxl import Workbook, load_workbook

from app.core.time_utils import local_now

DB = "postgresql://postgres:postgres@localhost:5432/bookera_test"


async def _biz(client, auth_headers, tag):
    h = auth_headers(f"io-owner-{tag}")
    bid = (await client.post("/crm/businesses", json={"name": "IO", "city": "Львів"}, headers=h)).json()["id"]
    return bid, h


def _xlsx(rows):
    wb = Workbook(); ws = wb.active
    for r in rows: ws.append(r)
    b = io.BytesIO(); wb.save(b); return b.getvalue()


@pytest.mark.asyncio
async def test_import_recognizes_columns_and_skips_duplicates(client, auth_headers):
    bid, h = await _biz(client, auth_headers, "imp")
    await client.post("/crm/clients", json={"business_id": bid, "name": "Вже є", "phone": "+380671111111"}, headers=h)
    data = _xlsx([
        ["ПІБ", "Номер", "E-mail", "Дата народження", "Теги"],
        ["Марія Коваль", "067 222 22 22", "Maria@Example.com", "14.03.1995", "VIP; постійна"],
        ["Ірина", "0501234567", None, None, None],
        ["Той самий номер", "+38 (067) 111-11-11", None, None, None],     # уже в базі
        ["Двічі у файлі", "067 222 22 22", None, None, None],             # повтор рядка 2
        ["Поганий номер", "12345", None, None, None],
        [None, "0673333333", None, None, None],                          # без імені
    ])
    files = {"file": ("база.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}

    dry = (await client.post("/crm/clients/import", params={"business_id": bid, "dry_run": "true"}, files=files, headers=h)).json()
    assert dry["to_create"] == 2 and dry["skipped"] == 4, dry
    reasons = {r["reason"] for r in dry["skipped_rows"]}
    assert {"уже є в базі", "номер не схожий на український", "немає імені"} <= reasons
    assert len((await client.get("/crm/clients", params={"business_id": bid}, headers=h)).json()) == 1, "перевірка нічого не створює"

    done = (await client.post("/crm/clients/import", params={"business_id": bid, "dry_run": "false"}, files=files, headers=h)).json()
    assert done["created"] == 2
    rows = {c["name"]: c for c in (await client.get("/crm/clients", params={"business_id": bid}, headers=h)).json()}
    assert rows["Марія Коваль"]["phone"] == "+380672222222" and rows["Марія Коваль"]["email"] == "maria@example.com"
    assert rows["Марія Коваль"]["birthday"] == "1995-03-14" and rows["Марія Коваль"]["tags"] == ["VIP", "постійна"]
    assert rows["Ірина"]["phone"] == "+380501234567"


@pytest.mark.asyncio
async def test_import_csv_cp1251_semicolon(client, auth_headers):
    """CSV з Excel українською: кодування Windows-1251 і крапка з комою."""
    bid, h = await _biz(client, auth_headers, "csv")
    data = "Імʼя;Телефон\nОлена;0679998877\n".replace("ʼ", "'").encode("cp1251")
    r = await client.post("/crm/clients/import", params={"business_id": bid, "dry_run": "false"},
                          files={"file": ("clients.csv", data, "text/csv")}, headers=h)
    assert r.status_code == 200 and r.json()["created"] == 1, r.text


@pytest.mark.asyncio
async def test_export_is_real_excel(client, auth_headers):
    bid, h = await _biz(client, auth_headers, "exp")
    await client.post("/crm/clients", json={"business_id": bid, "name": "Марія", "phone": "+380671234567"}, headers=h)
    r = await client.get("/crm/clients/export", params={"business_id": bid}, headers=h)
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A1"].value == "Імʼя" and ws["A2"].value == "Марія" and ws["B2"].value == "+380671234567"
    assert (await client.get("/crm/clients/export", params={"business_id": bid}, headers=auth_headers("io-stranger"))).status_code == 403


@pytest.mark.asyncio
async def test_merge_moves_everything(client, auth_headers, legacy_duplicates):
    """Обʼєднання переносить записи, сімейні звʼязки й бали; поля доповнюються."""
    bid, h = await _biz(client, auth_headers, "mrg")
    conn = await asyncpg.connect(DB)
    try:
        a = await conn.fetchval("INSERT INTO clients (business_id, name, phone, balance, tags, notes, is_blacklisted, consent_photo, consent_procedure) VALUES ($1,'Марія','+380671234567',100,'[\"VIP\"]','перша',false,false,false) RETURNING id", bid)
        b = await conn.fetchval("INSERT INTO clients (business_id, name, phone, email, birthday, balance, tags, notes, is_blacklisted, consent_photo, consent_procedure) VALUES ($1,'Маша','0671234567','m@example.com','1995-03-14',50,'[\"постійна\"]','друга',false,false,false) RETURNING id", bid)
        kid = await conn.fetchval("INSERT INTO clients (business_id, name, is_blacklisted, consent_photo, consent_procedure) VALUES ($1,'Донька',false,false,false) RETURNING id", bid)
        await conn.execute("INSERT INTO client_links (client_id, linked_client_id) VALUES ($1,$2)", b, kid)
        st = local_now().replace(tzinfo=None, microsecond=0) - timedelta(days=5)
        appt = await conn.fetchval("INSERT INTO appointments (business_id, client_id, client_name, start_time, end_time, status, price, source) VALUES ($1,$2,'Маша',$3,$4,'completed',500,'crm') RETURNING id", bid, b, st, st + timedelta(hours=1))
    finally:
        await conn.close()

    groups = (await client.get("/crm/clients/duplicates", params={"business_id": bid}, headers=h)).json()
    assert len(groups) == 1 and {x["id"] for x in groups[0]} == {a, b}

    r = await client.post(f"/crm/clients/{a}/merge", json={"merge_ids": [b]}, headers=h)
    assert r.status_code == 200, r.text
    rows = {c["id"]: c for c in (await client.get("/crm/clients", params={"business_id": bid}, headers=h)).json()}
    assert b not in rows
    m = rows[a]
    assert m["email"] == "m@example.com" and m["birthday"] == "1995-03-14", "порожні поля доповнено"
    assert float(m["balance"]) == 150 and set(m["tags"]) == {"VIP", "постійна"} and "друга" in m["notes"]
    assert m["visits_count"] == 1, "запис дубля - тепер у основній картці"
    assert kid in m["linked_client_ids"], "сімейний звʼязок перенесено"
    assert (await client.get("/crm/clients/duplicates", params={"business_id": bid}, headers=h)).json() == []


@pytest.mark.asyncio
async def test_template_is_empty_export_and_roundtrip(client, auth_headers):
    """Шаблон = експорт без даних; вивантажене можна завантажити назад."""
    bid, h = await _biz(client, auth_headers, "rt")
    tpl = load_workbook(io.BytesIO((await client.get("/crm/clients/import-template", headers=h)).content)).active
    exp_bid, h2 = await _biz(client, auth_headers, "rt2")
    await client.post("/crm/clients", json={"business_id": exp_bid, "name": "Марія", "phone": "+380671234567", "email": "m@example.com"}, headers=h2)
    exp = (await client.get("/crm/clients/export", params={"business_id": exp_bid}, headers=h2)).content
    headers_exp = [c.value for c in load_workbook(io.BytesIO(exp)).active[1]]
    assert [c.value for c in tpl[1]] == headers_exp and tpl.max_row == 1, "ті самі колонки, без даних"

    r = await client.post("/crm/clients/import", params={"business_id": bid, "dry_run": "false"},
                          files={"file": ("Клієнти.xlsx", exp, "application/octet-stream")}, headers=h)
    assert r.json()["created"] == 1
    got = (await client.get("/crm/clients", params={"business_id": bid}, headers=h)).json()[0]
    assert got["name"] == "Марія" and got["phone"] == "+380671234567" and got["email"] == "m@example.com"
