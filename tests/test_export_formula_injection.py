import io

import pytest


@pytest.mark.asyncio
async def test_export_does_not_turn_client_text_into_formulas(client, auth_headers):
    from openpyxl import load_workbook

    owner = auth_headers("xl-owner")
    bid = (await client.post("/crm/businesses", json={"name": "Excel", "city": "Львів"}, headers=owner)).json()["id"]
    evil = '=HYPERLINK("http://evil.example","клікни")'
    r = await client.post("/crm/clients", json={"business_id": bid, "name": evil, "phone": "+380671234567", "notes": "=1+1"}, headers=owner)
    assert r.status_code in (200, 201), r.text

    resp = await client.get(f"/crm/clients/export?business_id={bid}", headers=owner)
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    cells = [c for row in ws.iter_rows(min_row=2) for c in row if c.value]
    assert any(c.value == evil for c in cells), "текст збережено як є"
    assert all(c.data_type != "f" for c in cells), [(c.coordinate, c.data_type) for c in cells if c.data_type == "f"]
